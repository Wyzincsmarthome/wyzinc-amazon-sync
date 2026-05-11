# src/sources/visiotech.py
"""Read the Visiotech Connect XLSX catalog.

The file has no header row. Columns are positional. Confirmed mapping
(2026-05-11, by sampling the supplied file):

    col 1  -> SKU / reference
    col 3  -> stock bucket (text): "high" | "medium" | "low" | "none"
    col 5  -> brand
    col 9  -> short title
    col 13 -> PVD cost in EUR (what we pay Visiotech)
    col 16 -> EAN

Other columns (description HTML, images, weights) are not needed for a
price/stock update; they would only matter when creating new listings.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Iterable, Optional

import openpyxl

from src.core.models import ProductRecord
from src.sources.base import CatalogSource

logger = logging.getLogger(__name__)


# 1-based positions in the Visiotech sheet (matches the spreadsheet).
COL_SKU = 1
COL_STOCK_BUCKET = 3
COL_BRAND = 5
COL_TITLE = 9
COL_COST = 13
COL_EAN = 16

# Stock text -> advertised quantity on Amazon.
STOCK_BUCKET_QTY: dict[str, int] = {
    "high": 10,
    "medium": 5,
    "low": 1,
    "none": 0,
}


class VisiotechSource(CatalogSource):
    name = "visiotech"

    def __init__(
        self,
        path: str | Path,
        allowed_brands: Optional[list[str]] = None,
        sheet_name: str = "in",
    ) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"Visiotech file not found: {self.path}")
        self.sheet_name = sheet_name
        self.allowed_brands = {b.strip().lower() for b in (allowed_brands or []) if b}

    def fetch(self) -> Iterable[ProductRecord]:
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        if self.sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Sheet '{self.sheet_name}' not found in {self.path} "
                f"(available: {wb.sheetnames})"
            )
        sh = wb[self.sheet_name]

        kept = 0
        skipped_no_sku = 0
        skipped_brand = 0
        skipped_cost = 0
        skipped_stock_unknown = 0

        for row in sh.iter_rows(values_only=True):
            sku = _to_str(row[COL_SKU - 1])
            if not sku:
                skipped_no_sku += 1
                continue

            brand = _to_str(row[COL_BRAND - 1])
            if self.allowed_brands and (brand or "").lower() not in self.allowed_brands:
                skipped_brand += 1
                continue

            cost = _to_float(row[COL_COST - 1])
            if cost is None or cost <= 0:
                skipped_cost += 1
                continue

            stock_text = _to_str(row[COL_STOCK_BUCKET - 1]).lower()
            if stock_text not in STOCK_BUCKET_QTY:
                skipped_stock_unknown += 1
                continue
            quantity = STOCK_BUCKET_QTY[stock_text]

            ean = _to_str(row[COL_EAN - 1]) or None
            title = _to_str(row[COL_TITLE - 1]) or None

            kept += 1
            yield ProductRecord(
                sku=sku,
                cost=cost,
                stock=quantity,
                ean=ean,
                brand=brand or None,
                title=title,
                raw={"stock_bucket": stock_text},
            )

        wb.close()
        logger.info(
            "Visiotech parse: kept=%d, skipped_no_sku=%d, skipped_brand=%d, "
            "skipped_invalid_cost=%d, skipped_stock_unknown=%d",
            kept, skipped_no_sku, skipped_brand, skipped_cost, skipped_stock_unknown,
        )


def _to_str(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _to_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
